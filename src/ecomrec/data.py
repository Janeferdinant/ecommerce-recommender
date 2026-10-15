"""Load and clean UCI Online Retail II into a (user, item, ts) purchase log."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd

RAW_COLUMNS = ["Invoice", "StockCode", "Description", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country"]
# Real products have a 5-digit code with an optional letter suffix (e.g. 85123A).
# Everything else (POST, D, M, BANK CHARGES, gift vouchers, ...) is a fee or an adjustment.
PRODUCT_CODE = r"^\d{5}[A-Za-z]{0,2}$"


def read_workbook(xlsx_path: Path) -> pd.DataFrame:
    """Read both sheets of the workbook with openpyxl in streaming mode (a few minutes)."""
    from openpyxl import load_workbook

    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    frames = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        header = list(next(rows))
        frames.append(pd.DataFrame(list(rows), columns=header))
    workbook.close()
    return pd.concat(frames, ignore_index=True)[RAW_COLUMNS]


def clean_transactions(raw: pd.DataFrame) -> pd.DataFrame:
    """Keep completed purchases of real products by identified customers."""
    df = raw.copy()
    df["Invoice"] = df["Invoice"].astype(str)
    df["StockCode"] = df["StockCode"].astype(str).str.strip().str.upper()
    df = df[~df["Invoice"].str.upper().str.startswith("C")]  # cancellations / returns
    df = df[df["Customer ID"].notna()]
    df = df[(df["Quantity"] > 0) & (df["Price"] > 0)]
    df = df[df["StockCode"].str.match(PRODUCT_CODE, case=False)]
    df["InvoiceDate"] = pd.to_datetime(df["InvoiceDate"])
    return df.reset_index(drop=True)


def to_interactions(clean: pd.DataFrame) -> pd.DataFrame:
    """One row per (customer, product, basket): columns user, item, ts."""
    out = pd.DataFrame(
        {
            "user": clean["Customer ID"].astype(float).astype(int).astype(str),
            "item": clean["StockCode"],
            "ts": clean["InvoiceDate"],
            "invoice": clean["Invoice"],
            "price": clean["Price"].astype(float),
        }
    )
    out = out.drop_duplicates(["user", "item", "invoice"])
    return out[["user", "item", "ts", "price"]].reset_index(drop=True)


def item_names(clean: pd.DataFrame) -> dict:
    """Most frequent (title-cased) description per product code."""
    desc = clean.dropna(subset=["Description"]).copy()
    desc["Description"] = desc["Description"].astype(str).str.strip()
    top = (
        desc.groupby(["StockCode", "Description"]).size().reset_index(name="n")
        .sort_values(["StockCode", "n", "Description"], ascending=[True, False, True])
        .drop_duplicates("StockCode")
    )
    return {row.StockCode: str(row.Description).title() for row in top.itertuples()}


def load_clean(xlsx_path: Path, cache_dir: Path):
    """Return ``(interactions, names)``, caching the cleaned tables as CSV."""
    cache_dir.mkdir(parents=True, exist_ok=True)
    inter_path, names_path = cache_dir / "interactions.csv", cache_dir / "item_names.csv"
    if inter_path.exists() and names_path.exists():
        inter = pd.read_csv(inter_path, dtype={"user": str, "item": str}, parse_dates=["ts"])
        if "price" in inter.columns:  # caches written before the revenue proxy lack prices
            names = pd.read_csv(names_path, dtype=str).set_index("item")["name"].to_dict()
            return inter, names
    clean = clean_transactions(read_workbook(xlsx_path))
    inter, names = to_interactions(clean), item_names(clean)
    inter.to_csv(inter_path, index=False)
    pd.Series(names, name="name").rename_axis("item").reset_index().to_csv(names_path, index=False)
    return inter, names



# ---- Goodbooks-10k (Zygmunt Zajac, CC BY-SA 4.0) -----------------------------
# The ratings file carries no timestamps. Its documentation says the rows are sorted by time, so the position of a
# row in the file is used as a stand-in for time (one second per position). Only the order is meaningful.
GOODBOOKS_PSEUDO_START = pd.Timestamp("2000-01-01")


def read_goodbooks(ratings_zip: Path, books_zip: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Read ``ratings.csv`` and ``books.csv`` from the two downloaded release archives."""
    import io
    import zipfile

    with zipfile.ZipFile(ratings_zip) as archive:
        ratings = pd.read_csv(io.BytesIO(archive.read("ratings.csv")))
    with zipfile.ZipFile(books_zip) as archive:
        books = pd.read_csv(io.BytesIO(archive.read("books.csv")))
    return ratings[["user_id", "book_id", "rating"]], books


def goodbooks_interactions(ratings: pd.DataFrame, n_users: int | None = None, min_item_users: int = 1,
                           seed: int = 0) -> pd.DataFrame:
    """Every rating is one implicit interaction (columns user, item, ts, basket).

    ``n_users`` keeps a random sample of that many users (``seed``) with all their ratings, and books rated by fewer
    than ``min_item_users`` of the kept users are dropped. Row order, and therefore the stand-in time, is preserved.
    ``basket`` is the user id (there is no basket structure).
    """
    df = ratings.assign(position=np.arange(len(ratings)))
    if n_users is not None and n_users < df["user_id"].nunique():
        rng = np.random.default_rng(seed)
        chosen = rng.choice(np.sort(df["user_id"].unique()), size=n_users, replace=False)
        df = df[df["user_id"].isin(set(chosen.tolist()))]
    if min_item_users > 1:
        per_item = df.groupby("book_id")["user_id"].nunique()
        df = df[df["book_id"].isin(per_item[per_item >= min_item_users].index)]
    out = pd.DataFrame(
        {
            "user": df["user_id"].astype(int).astype(str).to_numpy(),
            "item": df["book_id"].astype(int).astype(str).to_numpy(),
            "ts": GOODBOOKS_PSEUDO_START + pd.to_timedelta(df["position"].to_numpy(), unit="s"),
        }
    )
    out["basket"] = out["user"]
    return out.drop_duplicates(["user", "item"]).reset_index(drop=True)


def goodbooks_item_info(books: pd.DataFrame) -> dict[str, str]:
    """Book titles keyed by item id (as strings)."""
    return {str(int(str(r.book_id))): str(r.title) for r in books.itertuples()}
